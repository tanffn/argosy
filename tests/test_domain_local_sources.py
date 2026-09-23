from datetime import date
import hashlib

from openpyxl import Workbook
import pytest

from argosy.services.domain_local_sources import read_local_record
from argosy.orchestrator.loops.annual import _attach_local_sources


@pytest.mark.asyncio
async def test_internal_knowledge_is_contained_bounded_and_dependency_versioned(engine, tmp_path, monkeypatch):
    from types import SimpleNamespace
    from argosy import config
    from argosy.orchestrator.loops import annual
    from argosy.services.knowledge_status import dependency_versions
    from argosy.state import db as db_mod
    root = tmp_path / 'knowledge'
    root.mkdir()
    sibling = root / 'sibling.md'
    sibling.write_text('# Internal context\nNot independent evidence', encoding='utf-8')
    (tmp_path / 'outside.md').write_text('private outside', encoding='utf-8')
    settings = SimpleNamespace(home=tmp_path, domain_knowledge_dir=root)
    monkeypatch.setattr(config, 'get_settings', lambda: settings)
    monkeypatch.setattr(annual, 'get_settings', lambda: settings)
    item = {'frontmatter': 'sources:\n - url: file://Knowledge/sibling.md\n - url: file://Knowledge/../outside.md'}
    prepared, attachments = _attach_local_sources(item)
    assert not attachments
    assert len(prepared['local_source_evidence']) == 1
    evidence = prepared['local_source_evidence'][0]
    assert evidence['status'] == 'captured' and 'not independent legal evidence' in evidence['provenance_note']
    assert 'private outside' not in str(prepared)
    async with db_mod.get_session() as session:
        before = await dependency_versions(session, user_id='ariel', item=item)
        sibling.write_text('changed sibling', encoding='utf-8')
        after = await dependency_versions(session, user_id='ariel', item=item)
        assert before != after
        assert ['file://Knowledge/../outside.md', 'outside_allowed_root'] in after


def test_markdown_catalog_source_is_bounded_and_original_hash_retained(tmp_path):
    path = tmp_path / 'historical.md'
    raw = ('# Original\n' + 'x' * 100_000).encode()
    path.write_bytes(raw)
    result = read_local_record(path, url='file://Catalog/121', as_of='2026-02 version')
    assert result['status'] == 'captured' and result['truncated']
    assert len(result['content']) == 100_000 and result['content'].startswith('# Original\n')
    assert result['sha256'] == hashlib.sha256(raw).hexdigest()


def test_html_as_xls_preserves_visible_headers_and_never_executes(tmp_path):
    path = tmp_path / 'wire.xls'
    raw = ('<HTML><style>hidden style</style><script>fetch("https://bad.invalid")</script>'
           '<body>חשבון פמ"ח<table><tr><th>Date</th><th>Credit</th></tr>'
           '<tr><td>12/08/2026</td><td>93,350.08</td></tr></table></body></HTML>').encode()
    path.write_bytes(raw)
    result = read_local_record(path, url='file://Resources/wire.xls', as_of='2026-08-22')
    assert result['status'] == 'captured' and result['format'] == 'html_as_xls'
    assert result['sha256'] == hashlib.sha256(raw).hexdigest()
    assert 'חשבון פמ"ח' in result['content']
    assert 'Date\nCredit\n12/08/2026\n93,350.08' in result['content']
    assert 'fetch(' not in result['content'] and 'hidden style' not in result['content']


@pytest.mark.parametrize('raw', [b'\xd0\xcf\x11\xe0binary', b'not HTML', b'<html>\xff</html>'])
def test_html_as_xls_rejects_binary_unrecognized_and_invalid_utf8(tmp_path, raw):
    path = tmp_path / 'wire.xls'
    path.write_bytes(raw)
    result = read_local_record(path, url='file://Resources/wire.xls')
    assert result['status'] == 'unavailable' and not result['content']


def _spreadsheetml(body):
    return ('\n <?xml version="1.0"?>'
            '<Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet" '
            'xmlns:ss="urn:schemas-microsoft-com:office:spreadsheet">'
            '<Worksheet ss:Name="תיק"><Table>' + body +
            '</Table></Worksheet></Workbook>').encode('utf-8')


def test_spreadsheetml_preserves_sparse_coordinates_types_and_unexecuted_formula(tmp_path):
    import json
    path = tmp_path / 'portfolio.xls'
    raw = _spreadsheetml('<Row ss:Index="3"><Cell ss:Index="2" ss:MergeAcross="2">'
                         '<Data ss:Type="String">שווי ₪</Data></Cell>'
                         '<Cell ss:Formula="=1+2"><Data ss:Type="Number">3</Data></Cell>'
                         '</Row><Row><Cell><Data ss:Type="DateTime">2026-09-11</Data></Cell></Row>')
    path.write_bytes(raw)
    result = read_local_record(path, url='file://Catalog/113', as_of='2026-09-11')
    assert result['status'] == 'captured' and result['format'] == 'spreadsheetml'
    assert result['sha256'] == hashlib.sha256(raw).hexdigest() and not result['truncated']
    assert result['source_as_of'] == '2026-09-11'
    first, second = [json.loads(line) for line in result['content'].splitlines()]
    assert first['sheet'] == 'תיק'
    assert first['cells']['B3'] == {'type': 'String', 'value': 'שווי ₪'}
    assert first['cells']['E3'] == {'type': 'Number', 'cached_value': '3', 'formula': '=1+2'}
    assert second['cells']['A4'] == {'type': 'DateTime', 'value': '2026-09-11'}


@pytest.mark.parametrize('bad', [
    b'<?xml version="1.0"?><Workbook/>',
    _spreadsheetml('<Row ss:Index="0"/>'),
    _spreadsheetml('<Row><Cell ss:Index="0"/></Row>'),
    _spreadsheetml('<Row><Cell/><Cell ss:Index="1"/></Row>'),
    _spreadsheetml('<Row><Cell ss:MergeAcross="-1"/></Row>'),
    _spreadsheetml('<Row/>').replace(b'<Table>', b'<Table ss:TopCell="2">'),
    _spreadsheetml('<Row/>').replace(b'<Workbook ',
        b'<!DOCTYPE Workbook [<!ENTITY x SYSTEM "file:///not-read">]><Workbook '),
    _spreadsheetml('<Row><Cell><Data>&unknown;</Data></Cell></Row>'),
    _spreadsheetml('<Row><Cell><Data>&é;</Data></Cell></Row>'),
])
def test_spreadsheetml_rejects_ambiguous_coordinates_external_entities_and_namespace(tmp_path, bad):
    path = tmp_path / 'bad.xls'
    path.write_bytes(bad)
    result = read_local_record(path, url='file://Resources/bad.xls')
    assert result['status'] == 'unavailable' and not result['content']
    assert result['sha256'] == hashlib.sha256(bad).hexdigest()


def test_spreadsheetml_huge_sparse_gaps_are_bounded_not_allocated(tmp_path):
    path = tmp_path / 'gap.xls'
    path.write_bytes(_spreadsheetml(
        '<Row><Cell><Data ss:Type="String">kept</Data></Cell>'
        '<Cell ss:Index="1000000000"><Data ss:Type="String">outside</Data></Cell></Row>'
        '<Row ss:Index="1000000000"><Cell><Data ss:Type="String">outside</Data></Cell></Row>'))
    result = read_local_record(path, url='file://Resources/gap.xls')
    assert result['status'] == 'captured' and result['truncated']
    assert 'kept' in result['content'] and 'outside' not in result['content']


def test_spreadsheetml_literal_ampersands_retained_without_loss_or_double_unescape(tmp_path):
    import json
    path = tmp_path / 'leumi.xls'
    path.write_bytes(_spreadsheetml('<Row><Cell><Data ss:Type="String">S&P 500</Data></Cell>'
        '<Cell><Data ss:Type="String">S&amp;P &amp;amp; &#38; &#x26;</Data></Cell></Row>'))
    result = read_local_record(path, url='file://Resources/leumi.xls')
    assert result['status'] == 'captured' and not result['truncated']
    assert result['normalizations'] == {'bare_ampersands_escaped': 1}
    cells = json.loads(result['content'])['cells']
    assert cells['A1']['value'] == 'S&P 500'
    assert cells['B1']['value'] == 'S&P &amp; & &'


def test_spreadsheetml_cdata_is_verbatim_and_comments_do_not_change_normalization(tmp_path):
    import json
    path = tmp_path / 'cdata.xls'
    path.write_bytes(_spreadsheetml('<!-- S&P --><?note S&P?>'
        '<Row><Cell><Data ss:Type="String"><![CDATA[S&P 500 &amp;]]></Data></Cell>'
        '<Cell><Data ss:Type="String">S&P 500</Data></Cell></Row>'))
    result = read_local_record(path, url='file://Resources/cdata.xls')
    assert result['status'] == 'captured' and not result['truncated']
    assert result['normalizations'] == {'bare_ampersands_escaped': 1}
    cells = json.loads(result['content'])['cells']
    assert cells['A1']['value'] == 'S&P 500 &amp;'
    assert cells['B1']['value'] == 'S&P 500'


def test_spreadsheetml_text_and_cell_budget_are_explicit(tmp_path, monkeypatch):
    from argosy.services import domain_local_sources as local
    path = tmp_path / 'large.xls'
    path.write_bytes(_spreadsheetml('<Row><Cell><Data ss:Type="String">first</Data></Cell></Row>'
        '<Row><Cell><Data ss:Type="String">' + 'x' * 60_001 + '</Data></Cell></Row>'))
    result = read_local_record(path, url='file://Resources/large.xls')
    assert result['status'] == 'captured' and result['truncated']
    assert 'first' in result['content'] and len(result['content']) < 60_000
    monkeypatch.setattr(local, 'MAX_CELLS', 1)
    path.write_bytes(_spreadsheetml('<Row><Cell><Data ss:Type="String">first</Data></Cell></Row>'
        '<Row><Cell><Data ss:Type="String">second</Data></Cell></Row>'))
    result = read_local_record(path, url='file://Resources/large.xls')
    assert result['truncated'] and 'second' not in result['content']


def test_html_as_xls_loader_keeps_limits_and_asof(tmp_path, monkeypatch):
    monkeypatch.setenv('ARGOSY_EXPENSE_SAMPLES_ROOT', str(tmp_path))
    (tmp_path / 'wire.xls').write_text('<html>' + 'x' * 60_001 + '</html>', encoding='utf-8')
    prepared, attachments = _attach_local_sources({'frontmatter': '''sources:
  - url: file://Resources/wire.xls
    as_of: 2026-08-22
'''})
    assert not attachments
    result = prepared['local_source_evidence'][0]
    assert result['status'] == 'captured' and result['truncated']
    assert len(result['content']) == 60_000 and result['source_as_of'] == '2026-08-22'


def test_workbook_preserves_coordinates_dates_and_formula_without_execution(tmp_path):
    path = tmp_path / 'record.xlsx'
    book = Workbook()
    book.active.title = 'Trustee'
    book.active.append(['Grant', date(2024, 4, 8), '=1+2', 18.1159])
    book.save(path)
    result = read_local_record(path, url='file://Resources/record.xlsx', as_of='2026-06-18')
    assert result['source_as_of'] == '2026-06-18'
    assert result['sha256'] == hashlib.sha256(path.read_bytes()).hexdigest()
    assert '"C1": "=1+2"' in result['content']
    assert 'Trustee' in result['content'] and '2024-04-08' in result['content']
    assert not result['truncated']


def test_text_bound_and_oversize_file(tmp_path):
    path = tmp_path / 'record.tsv'
    path.write_text('a' * 70_000, encoding='utf-8')
    result = read_local_record(path, url='file://Resources/record.tsv')
    assert result['truncated'] and len(result['content']) == 60_000
    assert result['source_as_of'] is None
    path.write_bytes(b'a' * 2_000_001)
    result = read_local_record(path, url='file://Resources/record.tsv')
    assert result['status'] == 'unavailable' and '2 MB' in result['error']
    assert 'sha256' not in result  # Partial raw bytes are not a complete-file hash.


def test_sparse_csv_retains_empty_duplicate_headers_continuations_and_formulas(tmp_path):
    import json
    path = tmp_path / 'awards.csv'
    raw = 'Date,Action,Value,Value,\r\n2026-02-06,Sale,520,,\r\n,,"multi\nline",=1+2,\r\n,,,,\r\n'
    path.write_bytes(raw.encode())
    result = read_local_record(path, url='file://Resources/awards.csv')
    assert result['status'] == 'captured' and not result['truncated']
    assert result['format'] == 'sparse_csv'
    lines = [json.loads(s) for s in result['content'].splitlines()]
    assert lines[0] == [1, 5, {'1': 'Date', '2': 'Action', '3': 'Value', '4': 'Value'}]
    assert lines[2] == [3, 5, {'3': 'multi\nline', '4': '=1+2'}]
    assert lines[3] == [4, 5, {}]


def test_sparse_csv_bounds_are_explicit(tmp_path):
    path = tmp_path / 'bounded.csv'
    path.write_text('first\n' + ','.join(['cell'] * 65) + '\n' * 2001, encoding='utf-8')
    result = read_local_record(path, url='file://Resources/bounded.csv')
    assert result['status'] == 'captured' and result['truncated']


def test_loader_refuses_traversal_and_preserves_historical_date(tmp_path, monkeypatch):
    root = tmp_path / 'Resources'
    root.mkdir()
    (root / 'snapshot.tsv').write_text('balance\t384000\n', encoding='utf-8')
    (tmp_path / 'outside.tsv').write_text('secret', encoding='utf-8')
    monkeypatch.setenv('ARGOSY_EXPENSE_SAMPLES_ROOT', str(root))
    prepared, attachments = _attach_local_sources({'frontmatter': '''sources:
  - url: file://Resources/snapshot.tsv
    as_of: 2025-12-31
  - url: file://Resources/%2e%2e/outside.tsv
'''})
    assert not attachments
    assert len(prepared['local_source_evidence']) == 1
    assert prepared['local_source_evidence'][0]['source_as_of'] == '2025-12-31'
    assert 'secret' not in str(prepared)
    assert 'NOT attached' in prepared['local_source_notes']


def test_loader_limits_aggregate_and_reports_omitted_sources(tmp_path, monkeypatch):
    (tmp_path / 'large.txt').write_text('a' * 70_000, encoding='utf-8')
    monkeypatch.setenv('ARGOSY_EXPENSE_SAMPLES_ROOT', str(tmp_path))
    prepared, _ = _attach_local_sources({'frontmatter': 'sources:\n' +
        '  - url: file://Resources/large.txt\n' * 4})
    assert sum(len(r['content']) for r in prepared['local_source_evidence']) == 180_000
    assert 'local record budget' in prepared['local_source_notes']
    assert all(r['truncated'] for r in prepared['local_source_evidence'])


def test_archive_expansion_and_wide_workbook_are_explicit(tmp_path):
    from zipfile import ZipFile, ZIP_DEFLATED
    path = tmp_path / 'record.xlsx'
    with ZipFile(path, 'w', compression=ZIP_DEFLATED) as archive:
        archive.writestr('oversize.xml', b'x' * 25_000_001)
    result = read_local_record(path, url='file://Resources/record.xlsx')
    assert result['status'] == 'unavailable' and 'expanded archive' in result['error']
    assert result['sha256']
    book = Workbook()
    book.active.cell(1, 65, 'outside bounded columns')
    book.save(path)
    result = read_local_record(path, url='file://Resources/record.xlsx')
    assert result['truncated']


def test_missing_dimensions_cannot_silently_omit_rows(tmp_path):
    import re
    from io import BytesIO
    from zipfile import ZipFile
    book = Workbook()
    book.active['A1'] = 'included'
    book.active['A2001'] = 'must not silently omit'
    original = BytesIO()
    book.save(original)
    path = tmp_path / 'no-dimensions.xlsx'
    with ZipFile(original) as before, ZipFile(path, 'w') as after:
        for entry in before.infolist():
            raw = before.read(entry.filename)
            if entry.filename == 'xl/worksheets/sheet1.xml':
                raw = re.sub(rb'<dimension[^>]*/>', b'', raw)
            after.writestr(entry, raw)
    result = read_local_record(path, url='file://Resources/no-dimensions.xlsx')
    assert result['status'] == 'captured' and result['truncated']
    assert 'included' in result['content']


def test_decode_failure_retains_provenance_in_loader(tmp_path, monkeypatch):
    path = tmp_path / 'bad.csv'
    path.write_bytes(b'\xff\xfe')
    monkeypatch.setenv('ARGOSY_EXPENSE_SAMPLES_ROOT', str(tmp_path))
    item, _ = _attach_local_sources({'frontmatter': '''sources:
  - url: file://Resources/bad.csv
    as_of: 2025-12-31
'''})
    receipt = item['local_source_evidence'][0]
    assert receipt['status'] == 'unavailable' and 'UnicodeDecodeError' in receipt['error']
    assert receipt['sha256'] == hashlib.sha256(path.read_bytes()).hexdigest()
    assert receipt['source_as_of'] == '2025-12-31' and receipt['accessed_at']


def test_understated_dimensions_do_not_hide_cells(tmp_path):
    import re
    from io import BytesIO
    from zipfile import ZipFile
    book = Workbook()
    book.active['A1'] = 'first'
    book.active['B2'] = 'must include'
    original = BytesIO()
    book.save(original)
    path = tmp_path / 'understated.xlsx'
    with ZipFile(original) as before, ZipFile(path, 'w') as after:
        for entry in before.infolist():
            raw = before.read(entry.filename)
            if entry.filename == 'xl/worksheets/sheet1.xml':
                raw = re.sub(rb'<dimension[^>]*/>', b'<dimension ref="A1:A1"/>', raw)
            after.writestr(entry, raw)
    result = read_local_record(path, url='file://Resources/understated.xlsx')
    assert result['status'] == 'captured' and not result['truncated']
    assert '"B2": "must include"' in result['content']


def test_implicit_coordinates_match_openpyxl(tmp_path):
    import re
    from io import BytesIO
    from zipfile import ZipFile
    book = Workbook()
    book.active['A1'] = 'first'
    book.active['B1'] = 'second'
    original = BytesIO()
    book.save(original)
    path = tmp_path / 'implicit.xlsx'
    with ZipFile(original) as before, ZipFile(path, 'w') as after:
        for entry in before.infolist():
            raw = before.read(entry.filename)
            if entry.filename == 'xl/worksheets/sheet1.xml':
                raw = re.sub(rb'(<c) r="[^"]*"', rb'\1', raw)
            after.writestr(entry, raw)
    result = read_local_record(path, url='file://Resources/implicit.xlsx')
    assert result['status'] == 'captured' and not result['truncated']
    assert '"A1": "first"' in result['content'] and '"B1": "second"' in result['content']


def test_prompt_distinguishes_attribution_and_optional_enrichment():
    from argosy.agents.domain_refresh import DomainRefreshAgent
    system, user = DomainRefreshAgent(user_id='test').build_prompt(files_due=[{
        'path': 'sample.md', 'local_source_evidence': [{'url': 'file://Resources/a.tsv', 'content': 'historical'}]}])
    assert 'not original provenance' in system
    assert 'Additions CAN be necessary repairs' in system
    assert 'Routine verification' in system
    assert 'current private-state assertions without current evidence remain partial' in system
    assert 'historical' in user and 'file://Resources/a.tsv' in user


@pytest.mark.asyncio
async def test_catalog_sources_enforce_owner_path_deletion_and_hash(engine, tmp_path, monkeypatch):
    from datetime import datetime
    from types import SimpleNamespace
    from argosy.orchestrator.loops import annual
    from argosy.state import db as db_mod
    from argosy.state.models import User, UserFile
    root = tmp_path / 'uploads' / 'ariel'
    root.mkdir(parents=True)
    good = root / 'record.txt'
    good.write_text('dated user statement', encoding='utf-8')
    outside = tmp_path / 'outside.txt'
    outside.write_text('never disclose', encoding='utf-8')
    monkeypatch.setattr(annual, 'get_settings', lambda: SimpleNamespace(home=tmp_path))
    async with db_mod.get_session() as session:
        session.add_all([User(id='ariel'), User(id='other')])
        for number in range(1, 6):
            session.add(UserFile(id=number, user_id='other' if number == 2 else 'ariel',
                sha256=hashlib.sha256(good.read_bytes()).hexdigest() if number == 1 else str(number) * 64,
                original_name='record.txt', sanitized_name='record.txt', mime_type='text/plain',
                kind='text', size_bytes=20, storage_path=str(outside if number == 4 else good),
                source='intake_file_to_text', deleted_at=datetime.now() if number == 3 else None))
        await session.commit()
    item = await annual._attach_catalog_sources({'frontmatter': 'catalog_sources:\n' +
        ''.join(f'  - id: {n}\n    as_of: 2026-08-23\n' for n in range(1, 6))}, user_id='ariel')
    evidence = item['local_source_evidence']
    assert evidence[0]['status'] == 'captured' and evidence[0]['source_as_of'] == '2026-08-23'
    assert all(e['status'] == 'unavailable' and not e['content'] for e in evidence[1:])
    assert 'never disclose' not in str(item)
    assert evidence[-1]['error'] == 'Catalog content hash mismatch'
