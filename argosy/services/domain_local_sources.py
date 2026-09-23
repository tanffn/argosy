"""Bounded read-only evidence from explicitly cited existing private records.

Caller resolves containment below the configured Resources directory. This does
not ingest, rewrite, execute formulas, or refresh the record's original date.
"""
from __future__ import annotations

import hashlib
import json
from datetime import UTC, date, datetime
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

MAX_BYTES = 2_000_000
MAX_TEXT = 60_000
MAX_CELLS = 10_000


def read_local_record(path: Path, *, url: str, as_of=None) -> dict:
    evidence = {'url': url, 'accessed_at': datetime.now(UTC).isoformat(),
                'source_as_of': str(as_of) if as_of is not None else None,
                'status': 'unavailable', 'content': '', 'truncated': False}
    try:
        with path.open('rb') as source:
            raw = source.read(MAX_BYTES + 1)
        if len(raw) > MAX_BYTES:
            raise ValueError('Cited local record exceeds 2 MB limit; complete bytes not read')
        evidence.update(sha256=hashlib.sha256(raw).hexdigest(), bytes=len(raw))
        _extract_record(path, raw, evidence)
        evidence['status'] = 'captured'
    except Exception as exc:
        evidence.update(error=f'{type(exc).__name__}: {exc}', content='')
    return evidence


def _extract_record(path: Path, raw: bytes, evidence: dict) -> dict:
    if path.suffix.lower() == '.xls':
        # Leumi exports HTML under an .xls name. This is evidence extraction,
        # NOT a cash-ledger import: retain headers/sub-account descriptions and
        # never merge cash and custody views. Binary XLS is not supported here.
        decoded = raw.decode('utf-8-sig')  # Refuse silent character replacement.
        prefix = decoded.lstrip().lower()
        if prefix.startswith('<?xml') or prefix.startswith('<workbook'):
            return _extract_spreadsheetml(decoded, evidence)
        if not (prefix.startswith('<html') or prefix.startswith('<!doctype html')):
            raise ValueError('Only UTF-8 HTML or SpreadsheetML XLS evidence is supported')
        from argosy.services.domain_sources import extract_text
        content = extract_text(raw, 'text/html')
        if not content.strip():
            raise ValueError('HTML-as-XLS record has no visible text')
        evidence['format'] = 'html_as_xls'
        evidence['truncated'] = len(content) > MAX_TEXT
        evidence['content'] = content[:MAX_TEXT]
        return evidence
    if path.suffix.lower() == '.csv':
        import csv
        from io import StringIO
        lines, length, count = [], 0, 0
        evidence['format'] = 'sparse_csv'
        evidence['layout'] = '[one-based row number, original column count, nonempty cells keyed by one-based column]; omitted cells are empty strings; first row is retained, no header inference'
        for number, row in enumerate(csv.reader(StringIO(raw.decode('utf-8-sig'), newline=''), strict=True), 1):
            if number > 2000 or len(row) > 64:
                evidence['truncated'] = True
            if number > 2000:
                continue
            cells = {str(index): value for index, value in enumerate(row[:64], 1) if value != ''}
            count += len(cells)
            line = json.dumps([number, len(row), cells], ensure_ascii=False, separators=(',', ':'))
            if count > MAX_CELLS or length + len(line) + 1 > MAX_TEXT:
                evidence['truncated'] = True
                break
            lines.append(line)
            length += len(line) + 1
        evidence['content'] = '\n'.join(lines)
        return evidence
    if path.suffix.lower() in ('.tsv', '.txt', '.md'):
        content = raw.decode('utf-8-sig')
        limit = 100_000 if path.suffix.lower() == '.md' else MAX_TEXT
        evidence['truncated'] = len(content) > limit
        evidence['content'] = content[:limit]
        return evidence
    if path.suffix.lower() != '.xlsx':
        raise ValueError('Unsupported local record format')
    with ZipFile(BytesIO(raw)) as archive:
        entries = archive.infolist()
        if len(entries) > 1000 or sum(e.file_size for e in entries) > 25_000_000:
            raise ValueError('Cited workbook exceeds expanded archive limit')
    from openpyxl import load_workbook
    from openpyxl.utils.cell import get_column_letter
    from openpyxl.worksheet._reader import WorkSheetParser
    book = load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    lines, length, cells = [], 0, 0
    try:
        for sheet in book.worksheets:
            # Use the installed openpyxl sparse parser that read-only worksheets
            # themselves use: actual/implicit coordinates, no declared dimensions
            # and no allocation of huge blank row/column gaps. Covered by contract
            # tests because this is an internal openpyxl integration.
            with sheet._get_source() as source:
                parser = WorkSheetParser(source, sheet._shared_strings, data_only=False,
                    epoch=book.epoch, date_formats=book._date_formats,
                    timedelta_formats=book._timedelta_formats)
                for row_number, row in parser.parse():
                    if row_number > 2000:
                        evidence['truncated'] = True
                        continue
                    values = {}
                    for cell in row:
                        if cell['row'] > 2000 or cell['column'] > 64:
                            evidence['truncated'] = True
                            continue
                        value = cell['value']
                        if value is None:
                            continue
                        if isinstance(value, (date, datetime)):
                            value = value.isoformat()
                        coordinate = f"{get_column_letter(cell['column'])}{cell['row']}"
                        values[coordinate] = value
                        cells += 1
                        if cells > MAX_CELLS:
                            evidence['truncated'] = True
                            evidence['content'] = '\n'.join(lines)
                            return evidence
                    if not values:
                        continue
                    line = json.dumps({'sheet': sheet.title, 'cells': values}, ensure_ascii=False, default=str)
                    if length + len(line) + 1 > MAX_TEXT:
                        evidence['truncated'] = True
                        evidence['content'] = '\n'.join(lines)
                        return evidence
                    lines.append(line)
                    length += len(line) + 1
    finally:
        book.close()
    evidence['content'] = '\n'.join(lines)
    return evidence


def _extract_spreadsheetml(decoded: str, evidence: dict) -> dict:
    """Read inert SpreadsheetML cells, not portfolio classifications or balances.

    Preserve sparse coordinates and distinguish formula text from cached output.
    Never allocate gaps or trust ExpandedRow/ColumnCount; same bounds as XLSX.
    """
    from lxml import etree
    from openpyxl.utils.cell import get_column_letter
    import re

    if '<!DOCTYPE' in decoded.upper() or '<!ENTITY' in decoded.upper():
        raise ValueError('DTD/entity declarations are not supported in spreadsheet evidence')
    # Actual Leumi exports contain literal "S&P" inside Data elements. Escape
    # only bare ampersands; preserve declared references and reject unknown ones.
    # Never use XML recover=True, which can silently discard source characters.
    parts = re.split(r'(<!\[CDATA\[.*?\]\]>|<!--.*?-->|<\?.*?\?>)', decoded, flags=re.S)
    escaped = 0
    for index in range(0, len(parts), 2):
        # Leave every entity-shaped token to the strict XML parser, including
        # Unicode names. CDATA/comments/processing instructions stay verbatim.
        parts[index], count = re.subn(r'&(?![^\s<&;]*;)', '&amp;', parts[index])
        escaped += count
    decoded = ''.join(parts)
    if escaped:
        evidence['normalizations'] = {'bare_ampersands_escaped': escaped}
    parser = etree.XMLParser(resolve_entities=False, load_dtd=False, no_network=True)
    root = etree.fromstring(decoded.lstrip().encode('utf-8'), parser=parser)
    ns = '{urn:schemas-microsoft-com:office:spreadsheet}'
    if root.tag != ns + 'Workbook':
        raise ValueError('Expected SpreadsheetML Workbook namespace')
    lines, length, cells = [], 0, 0
    evidence['format'] = 'spreadsheetml'
    for sheet in root.findall(ns + 'Worksheet'):
        for table in sheet.findall(ns + 'Table'):
            if any(int(table.get(ns + name, 1)) != 1 for name in ('TopCell', 'LeftCell')):
                raise ValueError('SpreadsheetML table offsets are not supported')
            row_number = 0
            for row in table.findall(ns + 'Row'):
                row_index = int(row.get(ns + 'Index', row_number + 1))
                if row_index <= row_number:
                    raise ValueError('SpreadsheetML row index is not increasing')
                row_number = row_index
                values, column = {}, 1
                for cell in row.findall(ns + 'Cell'):
                    index = int(cell.get(ns + 'Index', column))
                    across = int(cell.get(ns + 'MergeAcross', 0))
                    if index < column or across < 0:
                        raise ValueError('Invalid SpreadsheetML cell index or merge')
                    column = index + across + 1
                    if row_number > 2000 or index > 64:
                        evidence['truncated'] = True
                        continue
                    data = cell.find(ns + 'Data')
                    formula = cell.get(ns + 'Formula')
                    if data is None and formula is None:
                        continue
                    value = ''.join(data.itertext()) if data is not None else None
                    entry = {'type': data.get(ns + 'Type') if data is not None else None}
                    entry['cached_value' if formula is not None else 'value'] = value
                    if formula is not None:
                        entry['formula'] = formula
                    values[f'{get_column_letter(index)}{row_number}'] = entry
                    cells += 1
                    if cells > MAX_CELLS:
                        evidence.update(truncated=True, content='\n'.join(lines))
                        return evidence
                if values:
                    line = json.dumps({'sheet': sheet.get(ns + 'Name'), 'cells': values}, ensure_ascii=False)
                    if length + len(line) + 1 > MAX_TEXT:
                        evidence.update(truncated=True, content='\n'.join(lines))
                        return evidence
                    lines.append(line)
                    length += len(line) + 1
    if not lines:
        raise ValueError('SpreadsheetML evidence has no readable cells within bounds')
    evidence['content'] = '\n'.join(lines)
    return evidence
